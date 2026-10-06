"""An auditable Excel review pack generated from one immutable run."""
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _cell(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def _sheet(workbook: Workbook, name: str, headers: list[str], rows: list[list[Any]], money: tuple[int, ...] = ()) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell(value) for value in row])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='153B3B')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Calibri', size=11, color='253A3D')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F1F6F3')
        for column in money:
            row[column - 1].number_format = '"₪" #,##0.00;[Red]("₪" #,##0.00)'
    for column, header in enumerate(headers, 1):
        sample_lengths = [len(str(sheet.cell(row, column).value or '')) for row in range(1, min(sheet.max_row, 100) + 1)]
        sheet.column_dimensions[get_column_letter(column)].width = min(60, max(16, max(sample_lengths, default=len(header)) + 2))


def build_report(run: dict) -> bytes:
    """Return a formula-safe workbook. Amounts are exported in ILS, not agorot."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    audit = run.get('run', run)
    metrics = run.get('metrics', {})
    _sheet(workbook, 'Summary', ['Control', 'Value'], [
        ['Dataset', 'SYNTHETIC DATA — independent portfolio, no Harel affiliation'],
        ['Purpose', 'Reconciliation review only — no salary payment or statutory payroll'],
        ['Run ID', audit['id']], ['Period', audit['period']], ['Status', audit['status']],
        ['Publication eligibility', 'BLOCKED' if audit['critical_count'] else 'Passed critical controls; warnings still need review'],
        ['Source SHA256', audit['source_hash']], ['Rule version', run.get('rule_version', '1')],
        ['Started at (UTC)', audit['started_at']], ['Accepted rows', audit['accepted_rows']],
        ['Rejected source rows', audit['rejected_rows']], ['Critical findings', audit['critical_count']],
        ['Warning findings', audit['warning_count']], ['Headcount', metrics.get('headcount', 0)],
        ['Gross (ILS)', metrics.get('gross_cents', 0) / 100],
        ['Expected gross (ILS)', metrics.get('expected_cents', 0) / 100],
        ['Variance (ILS)', metrics.get('variance_cents', 0) / 100],
        ['Currency conversion', 'Money stored as integer agorot; displayed here in ILS'],
    ])
    for row in (16, 17, 18):
        workbook['Summary'].cell(row, 2).number_format = '"₪" #,##0.00'
    _sheet(workbook, 'Reconciliation', ['Employee ID', 'Synthetic employee', 'Department', 'Expected ILS', 'Reported ILS', 'Variance ILS', 'Overtime hours', 'Review status'], [
        [r['employee_id'], r['employee_name'], r['department'], r['expected_cents'] / 100,
         r['gross_cents'] / 100, r['variance_cents'] / 100, r['overtime_minutes'] / 60, r['status']]
        for r in run.get('records', [])
    ], money=(4, 5, 6))
    _sheet(workbook, 'Exceptions', ['Finding ID', 'Employee ID', 'Synthetic employee', 'Department', 'Rule', 'Severity', 'Evidence', 'Owner', 'Remediation', 'Impact ILS'], [
        [f['id'], f['employee_id'], f['employee_name'], f['department'], f['rule'], f['severity'], f['message'], f['owner'], f['remediation'], f.get('impact_cents', 0) / 100]
        for f in run.get('findings', [])
    ], money=(10,))
    _sheet(workbook, 'Lineage', ['Source', 'Interface', 'Rows', 'Status', 'Source file SHA256', 'Snapshot path', 'Run source SHA256'], [
        [s['name'], s['type'], s['rows'], s['status'], s.get('sha256', 'Not recorded'), s.get('path', 'Not recorded'), audit['source_hash']] for s in run.get('sources', [])
    ])
    workbook.properties.title = f"PeopleOps review {audit['period']}"
    workbook.properties.description = 'Synthetic, immutable ETL evidence. Independent portfolio.'
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
