"""Excel evidence contracts exercised through real immutable pipeline runs."""
from io import BytesIO
import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from peopleops.fixtures import employee_fixture
from peopleops.pipeline import PeopleOpsService
from peopleops.reports import build_report


class ReportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.data_dir = Path(self.directory.name)
        self.service = PeopleOpsService(self.data_dir)

    def workbook(self, evidence: dict) -> Workbook:
        return load_workbook(BytesIO(build_report(evidence)), data_only=False)

    @staticmethod
    def summary(workbook: Workbook) -> dict:
        return dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))

    def test_blocked_report_exposes_quarantine_counts_and_financial_worked_examples(self) -> None:
        workbook = self.workbook(self.service.run("2026-09", "review"))
        self.assertEqual(workbook.sheetnames, ["Summary", "Reconciliation", "Exceptions", "Lineage"])
        summary = self.summary(workbook)
        self.assertIn("SYNTHETIC DATA", summary["Dataset"])
        self.assertIn("no salary payment", summary["Purpose"])
        self.assertEqual(summary["Period"], "2026-09")
        self.assertEqual(summary["Status"], "blocked")
        self.assertEqual(summary["Publication eligibility"], "BLOCKED")
        self.assertEqual(summary["Accepted rows"], 211)
        self.assertEqual(summary["Rejected source rows"], 5)
        self.assertEqual(summary["Critical findings"], 7)
        self.assertEqual(summary["Warning findings"], 1)
        self.assertEqual(summary["Headcount"], 67)
        self.assertEqual(summary["Variance (ILS)"], 1865)
        self.assertEqual(workbook["Reconciliation"].max_row, 68)
        self.assertEqual(workbook["Exceptions"].max_row, 9)

        records = {row[0]: row for row in workbook["Reconciliation"].iter_rows(min_row=2, values_only=True)}
        # EMP-001: ILS 9,731 + 6.5 hours * ILS 52.32 * 1.25 = ILS 10,156.10.
        self.assertEqual(records["EMP-001"][3:8], (10156.10, 10156.10, 0, 6.5, "matched"))
        # EMP-007: ILS 14,117 + 8.5 * ILS 75.90 * 1.25 = 14,923.4375,
        # rounded half up to 14,923.44; the faulty gross has an extra ILS 1,865.
        self.assertEqual(records["EMP-007"][3:8], (14923.44, 16788.44, 1865, 8.5, "critical"))
        for employee_id in ("EMP-004", "EMP-010", "EMP-018", "EMP-029", "EMP-072"):
            self.assertNotIn(employee_id, records)

    def test_exception_pack_routes_remediation_and_formats_values_for_review(self) -> None:
        workbook = self.workbook(self.service.run("2026-09", "review"))
        findings = {row[4]: row for row in workbook["Exceptions"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(findings["duplicate_source_key"][7], "Data Integration")
        self.assertIn("Resolve duplicate payroll keys", findings["duplicate_source_key"][8])
        self.assertEqual(findings["missing_attendance"][7], "Time & Attendance")
        self.assertIn("2026-09", findings["missing_attendance"][8])
        self.assertEqual(findings["gross_difference"][7], "Payroll")
        self.assertIn("rerun", findings["gross_difference"][8])
        self.assertEqual(findings["gross_difference"][9], 1865)
        self.assertEqual(findings["overtime_threshold"][5], "warning")
        self.assertEqual(findings["overtime_threshold"][7], "Department Manager")
        for row in findings.values():
            self.assertTrue(row[7], "Every exception needs an owner")
            self.assertTrue(row[8], "Every exception needs a corrective action")

        for sheet in workbook:
            with self.subTest(sheet=sheet.title):
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        summary_cells = {row[0].value: row[1] for row in workbook["Summary"].iter_rows(min_row=2)}
        for label in ("Gross (ILS)", "Expected gross (ILS)", "Variance (ILS)"):
            self.assertIn("₪", summary_cells[label].number_format)
        for label in ("Headcount", "Accepted rows", "Rejected source rows"):
            self.assertNotIn("₪", summary_cells[label].number_format)
        for column in (4, 5, 6):
            self.assertIn("₪", workbook["Reconciliation"].cell(2, column).number_format)
        self.assertNotIn("₪", workbook["Reconciliation"].cell(2, 7).number_format)
        self.assertIn("₪", workbook["Exceptions"].cell(2, 10).number_format)

    def test_lineage_identifies_each_retained_file_and_its_run_checksum(self) -> None:
        evidence = self.service.run("2026-09", "review")
        workbook = self.workbook(evidence)
        summary = self.summary(workbook)
        self.assertEqual(summary["Run ID"], evidence["run"]["id"])
        self.assertEqual(summary["Rule version"], "simplified-gross-v1")
        self.assertEqual(summary["Source SHA256"], evidence["run"]["source_hash"])
        self.assertRegex(summary["Source SHA256"], r"^[a-f0-9]{64}$")
        lineage = {row[0]: row for row in workbook["Lineage"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(set(lineage), {"employees", "attendance", "payroll"})
        for name, row in lineage.items():
            with self.subTest(source=name):
                self.assertEqual(row[2], 72)
                self.assertEqual(row[3], "validated" if name == "employees" else "quarantined")
                self.assertRegex(row[4], r"^[a-f0-9]{64}$")
                snapshot = self.data_dir / row[5]
                self.assertTrue(snapshot.is_file())
                self.assertEqual(hashlib.sha256(snapshot.read_bytes()).hexdigest(), row[4])
                self.assertEqual(row[6], summary["Source SHA256"])
        self.assertEqual(lineage["employees"][1], "Synthetic HR snapshot (JSON)")
        self.assertEqual(lineage["attendance"][1], "Synthetic CSV export")
        self.assertEqual(lineage["payroll"][1], "Synthetic CSV export")

    def test_historical_blocked_export_stays_immutable_after_corrected_publication(self) -> None:
        blocked = self.service.run("2026-09", "review")
        original_book = self.workbook(blocked)
        original_values = {sheet.title: list(sheet.values) for sheet in original_book}

        corrected = self.service.run("2026-09", "corrected")
        current_book = self.workbook(self.service.get_run(corrected["run"]["id"]))
        summary = self.summary(current_book)
        self.assertEqual(summary["Status"], "published")
        self.assertIn("Passed critical controls", summary["Publication eligibility"])
        self.assertEqual(summary["Accepted rows"], 216)
        self.assertEqual(summary["Rejected source rows"], 0)
        self.assertEqual(summary["Critical findings"], 0)
        self.assertEqual(summary["Warning findings"], 0)
        self.assertEqual(summary["Headcount"], 72)
        self.assertEqual(summary["Variance (ILS)"], 0)
        self.assertEqual(current_book["Reconciliation"].max_row, 73)
        self.assertEqual(current_book["Exceptions"].max_row, 1)
        self.assertEqual(current_book["Exceptions"].auto_filter.ref, "A1:J1")
        self.assertNotEqual(summary["Run ID"], self.summary(original_book)["Run ID"])
        self.assertNotEqual(summary["Source SHA256"], self.summary(original_book)["Source SHA256"])

        historical = self.service.get_run(blocked["run"]["id"])
        historical_book = self.workbook(historical)
        self.assertEqual({sheet.title: list(sheet.values) for sheet in historical_book}, original_values)
        self.assertEqual(self.service.dashboard("2026-09")["metrics"]["published_run_id"], corrected["run"]["id"])
        for source in historical_book["Lineage"].iter_rows(min_row=2, values_only=True):
            self.assertEqual(hashlib.sha256((self.data_dir / source[5]).read_bytes()).hexdigest(), source[4])

    def test_external_hr_names_export_as_text_for_every_formula_prefix(self) -> None:
        # Supply actual HR source rows. EMP-007 appears in both the reconciled
        # records and the gross-difference exceptions in the review scenario.
        for name in ("=1+1", "+SUM(A1:A2)", "-1+1", "@SUM(A1:A2)",
                     "  =1+1", "\t+SUM(A1:A2)", " \t-1+1", "  @SUM(A1:A2)"):
            with self.subTest(external_name=name):
                employees = employee_fixture()
                employees[6]["employee_name"] = name
                employees[6]["department"] = name
                evidence = self.service.run("2026-09", "review", employees=employees)
                workbook = self.workbook(self.service.get_run(evidence["run"]["id"]))
                for sheet_name in ("Reconciliation", "Exceptions"):
                    sheet = workbook[sheet_name]
                    id_column = 1 if sheet_name == "Reconciliation" else 2
                    name_column = 2 if sheet_name == "Reconciliation" else 3
                    employee_row = next(row for row in sheet.iter_rows(min_row=2)
                                        if row[id_column - 1].value == "EMP-007")
                    for cell in employee_row[name_column - 1:name_column + 1]:
                        self.assertEqual(cell.value, "'" + name)
                        self.assertEqual(cell.data_type, "s")
                for sheet in workbook:
                    self.assertTrue(all(cell.data_type != "f" for row in sheet for cell in row),
                                    "Source text must never create an Excel formula")
                hr_lineage = next(row for row in workbook["Lineage"].iter_rows(min_row=2, values_only=True)
                                  if row[0] == "employees")
                self.assertEqual(hr_lineage[1], "HR API snapshot (JSON)")


if __name__ == "__main__":
    unittest.main()
