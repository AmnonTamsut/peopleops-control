"""HTTP integration contracts exercised with real synthetic ETL and storage."""
import asyncio
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from typing import Any

import httpx
from openpyxl import load_workbook

from peopleops.api import create_app


class ApiTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.app = create_app(Path(self.directory.name), seed=False)
        self.lifespan = self.app.router.lifespan_context(self.app)
        await self.lifespan.__aenter__()
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.app), base_url="http://test.local"
        )
        # The managed sandbox scheduler can miss AnyIO worker wakeups. A small
        # test-only heartbeat keeps the event loop responsive during HTTP calls.
        async def heartbeat() -> None:
            while True:
                await asyncio.sleep(0.01)

        self.heartbeat = asyncio.create_task(heartbeat())

    async def asyncTearDown(self) -> None:
        await self.client.aclose()
        await self.lifespan.__aexit__(None, None, None)
        self.heartbeat.cancel()
        try:
            await self.heartbeat
        except asyncio.CancelledError:
            pass

    async def request(self, method: str, url: str, **kwargs: Any) -> httpx.Response:
        return await asyncio.wait_for(self.client.request(method, url, **kwargs), timeout=5)

    async def run_scenario(self, scenario: str, period: str = "2026-09") -> dict:
        response = await self.request("POST", "/api/runs", json={"period": period, "scenario": scenario})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    async def get_json(self, url: str, **kwargs: Any) -> dict:
        response = await self.request("GET", url, **kwargs)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    async def test_publication_gate_keeps_last_good_data_and_preserves_historical_runs(self) -> None:
        clean = await self.run_scenario("clean")
        self.assertEqual(clean["run"]["status"], "published")
        self.assertEqual(clean["metrics"]["headcount"], 72)
        self.assertEqual(clean["metrics"]["published_run_id"], clean["run"]["id"])
        saved_clean = await self.get_json(f'/api/runs/{clean["run"]["id"]}')

        review = await self.run_scenario("review")
        self.assertEqual(review["run"]["status"], "blocked")
        self.assertEqual(review["run"]["critical_count"], 7)
        self.assertEqual(review["run"]["warning_count"], 1)
        self.assertEqual(review["run"]["rejected_rows"], 5)
        self.assertEqual(review["metrics"]["headcount"], 67)
        dashboard = await self.get_json("/api/dashboard", params={"period": "2026-09"})
        self.assertEqual(dashboard["run"]["id"], review["run"]["id"])
        self.assertEqual(dashboard["metrics"]["published_run_id"], clean["run"]["id"])
        saved_review = await self.get_json(f'/api/runs/{review["run"]["id"]}')

        corrected = await self.run_scenario("corrected")
        self.assertEqual(corrected["run"]["status"], "published")
        self.assertEqual(corrected["run"]["critical_count"], 0)
        self.assertEqual(corrected["metrics"]["headcount"], 72)
        dashboard = await self.get_json("/api/dashboard")
        self.assertEqual(dashboard["metrics"]["published_run_id"], corrected["run"]["id"])
        self.assertEqual(len(dashboard["runs"]), 3)
        self.assertEqual(await self.get_json(f'/api/runs/{clean["run"]["id"]}'), saved_clean)
        self.assertEqual(await self.get_json(f'/api/runs/{review["run"]["id"]}'), saved_review)

    async def test_identical_http_requests_reuse_evidence_without_duplicate_history(self) -> None:
        first = await self.run_scenario("corrected")
        saved = await self.get_json(f'/api/runs/{first["run"]["id"]}')
        history = (await self.get_json("/api/dashboard"))["runs"]

        repeat = await self.run_scenario("corrected")

        self.assertTrue(repeat["run"]["reused"])
        self.assertEqual(repeat["run"]["id"], first["run"]["id"])
        self.assertEqual(repeat["records"], first["records"])
        self.assertEqual((await self.get_json("/api/dashboard"))["runs"], history)
        self.assertEqual(await self.get_json(f'/api/runs/{first["run"]["id"]}'), saved)

    async def test_period_filter_isolates_runs_and_publication_pointer(self) -> None:
        september = await self.run_scenario("clean", "2026-09")
        october = await self.run_scenario("review", "2026-10")

        september_dashboard = await self.get_json("/api/dashboard", params={"period": "2026-09"})
        self.assertEqual(september_dashboard["run"]["id"], september["run"]["id"])
        self.assertEqual(september_dashboard["metrics"]["published_run_id"], september["run"]["id"])
        self.assertEqual(len(september_dashboard["runs"]), 1)
        october_dashboard = await self.get_json("/api/dashboard", params={"period": "2026-10"})
        self.assertEqual(october_dashboard["run"]["id"], october["run"]["id"])
        self.assertIsNone(october_dashboard["metrics"]["published_run_id"])
        self.assertEqual(len(october_dashboard["runs"]), 1)
        empty = await self.get_json("/api/dashboard", params={"period": "2026-11"})
        self.assertIsNone(empty["run"])
        self.assertEqual(empty["runs"], [])
        self.assertEqual(empty["metrics"]["headcount"], 0)

    async def test_report_download_is_a_readable_historical_excel_review_pack(self) -> None:
        blocked = await self.run_scenario("review")
        run_id = blocked["run"]["id"]
        url = f"/api/runs/{run_id}/report"
        response = await self.request("GET", url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(response.headers["content-disposition"],
                         f'attachment; filename="peopleops-2026-09-{run_id}.xlsx"')
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Summary", "Reconciliation", "Exceptions", "Lineage"])
        summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["Run ID"], run_id)
        self.assertEqual(summary["Status"], "blocked")
        self.assertEqual(summary["Publication eligibility"], "BLOCKED")
        self.assertEqual(summary["Headcount"], 67)
        self.assertEqual(summary["Variance (ILS)"], 1865)
        self.assertEqual(workbook["Reconciliation"].max_row, 68)
        self.assertEqual(workbook["Exceptions"].max_row, 9)
        original_values = {sheet.title: list(sheet.values) for sheet in workbook}

        await self.run_scenario("corrected")
        historical_response = await self.request("GET", url)
        self.assertEqual(historical_response.status_code, 200)
        historical_workbook = load_workbook(BytesIO(historical_response.content), data_only=False)
        self.addCleanup(historical_workbook.close)
        self.assertEqual({sheet.title: list(sheet.values) for sheet in historical_workbook}, original_values)

    async def test_unknown_run_and_report_return_404(self) -> None:
        for suffix in ("", "/report"):
            with self.subTest(suffix=suffix):
                response = await self.request("GET", f"/api/runs/not-a-known-run{suffix}")
                self.assertEqual(response.status_code, 404)
                self.assertEqual(response.json(), {"detail": "Run not found"})

    async def test_invalid_run_input_is_rejected_without_creating_audit_entries(self) -> None:
        invalid_bodies = (
            {"period": "2026-13", "scenario": "clean"},
            {"period": "2026-00", "scenario": "clean"},
            {"period": "2026-9", "scenario": "clean"},
            {"period": "0000-01", "scenario": "clean"},
            {"period": "2026-09", "scenario": "unrecognized"},
            {"period": None, "scenario": "clean"},
        )
        for body in invalid_bodies:
            with self.subTest(body=body):
                response = await self.request("POST", "/api/runs", json=body)
                self.assertEqual(response.status_code, 422, response.text)
        dashboard = await self.get_json("/api/dashboard")
        self.assertIsNone(dashboard["run"])
        self.assertEqual(dashboard["runs"], [])

    async def test_invalid_dashboard_period_returns_422(self) -> None:
        for period in ("2026-13", "2026-00", "2026-9", "garbage", "0000-01"):
            with self.subTest(period=period):
                response = await self.request("GET", "/api/dashboard", params={"period": period})
                self.assertEqual(response.status_code, 422, response.text)

    async def test_mock_hr_pagination_returns_all_72_employees_once(self) -> None:
        employees = []
        for page, size, has_more in ((1, 25, True), (2, 25, True), (3, 22, False), (4, 0, False)):
            response = await self.get_json("/mock/hr/employees", params={"page": page, "page_size": 25})
            self.assertEqual(len(response["items"]), size)
            self.assertEqual(response["has_more"], has_more)
            employees.extend(response["items"])
        self.assertEqual(len(employees), 72)
        self.assertEqual(len({row["employee_id"] for row in employees}), 72)
        self.assertEqual({row["employee_id"] for row in employees},
                         {f"EMP-{number:03}" for number in range(1, 73)})
        for employee in employees:
            self.assertIs(type(employee["active"]), bool)
            self.assertIs(type(employee["base_salary_cents"]), int)
            self.assertIs(type(employee["hourly_cents"]), int)
        full_page = await self.get_json("/mock/hr/employees", params={"page_size": 100})
        self.assertEqual(full_page["items"], employees)
        self.assertFalse(full_page["has_more"])

    async def test_mock_hr_invalid_pagination_is_rejected(self) -> None:
        for params in ({"page": 0}, {"page": -1}, {"page": "garbage"},
                       {"page_size": 0}, {"page_size": 101}, {"page_size": "garbage"}):
            with self.subTest(params=params):
                response = await self.request("GET", "/mock/hr/employees", params=params)
                self.assertEqual(response.status_code, 422, response.text)

    async def test_health_and_browser_assets_are_available_without_running_pipeline(self) -> None:
        health = await self.get_json("/api/health")
        self.assertEqual(health, {"status": "ok", "storage": "sqlite", "dataset": "synthetic", "mode": "local-demo"})
        for url, media_type, expected_text in (
            ("/", "text/html", "PeopleOps"),
            ("/static/app.js", "javascript", "/api/dashboard"),
            ("/static/styles.css", "text/css", "body"),
        ):
            with self.subTest(url=url):
                response = await self.request("GET", url)
                self.assertEqual(response.status_code, 200)
                self.assertIn(media_type, response.headers["content-type"])
                self.assertIn(expected_text, response.text)
        dashboard = await self.get_json("/api/dashboard")
        self.assertIsNone(dashboard["run"])
        self.assertEqual(dashboard["runs"], [])

    async def test_default_lifespan_seeds_demo_once_and_retains_last_good_pointer(self) -> None:
        seeded_app = create_app(Path(self.directory.name) / "seeded", seed=True)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=seeded_app),
                                     base_url="http://test.local") as client:
            for startup in range(2):
                async with seeded_app.router.lifespan_context(seeded_app):
                    response = await asyncio.wait_for(client.get("/api/dashboard"), timeout=5)
                    self.assertEqual(response.status_code, 200)
                    dashboard = response.json()
                    self.assertEqual(len(dashboard["runs"]), 4)
                    self.assertEqual(dashboard["run"]["scenario"], "review")
                    self.assertEqual(dashboard["run"]["status"], "blocked")
                    published_id = dashboard["metrics"]["published_run_id"]
                    self.assertIsNotNone(published_id)
                    response = await asyncio.wait_for(client.get(f"/api/runs/{published_id}"), timeout=5)
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(response.json()["run"]["scenario"], "clean")
                    if startup == 0:
                        first_dashboard = dashboard
                    else:
                        self.assertEqual(dashboard, first_dashboard)


if __name__ == "__main__":
    unittest.main()
